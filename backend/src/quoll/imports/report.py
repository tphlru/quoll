"""Итог партии в xlsx (import-design §21): лист на лист файла - номер строки,
статус, итог, замечания и исходные ячейки сопоставленных колонок"""

import io
import re

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from quoll.imports.models import ImportRow, ImportSheet
from quoll.imports.spec import ISSUES, KINDS

STATUSES = {
    "NEW": "новая",
    "UPDATE": "изменится",
    "SAME": "без изменений",
    "CONFLICT": "конфликт",
    "ERROR": "ошибка",
    "EXCLUDED": "исключена",
}
RESULTS = {"APPLIED": "применена", "SKIPPED": "пропущена", "FAILED": "не удалась"}


def _title(sheet: ImportSheet, used: set[str]) -> str:
    name = re.sub(r"[\[\]:*?/\\]", " ", f"{sheet.number}. {sheet.name}")[:31]
    while name in used:
        name = name[:28] + f"~{len(used)}"
    used.add(name)
    return name


def _issues(row: ImportRow) -> str:
    return "; ".join(
        f"{i['code']} {ISSUES.get(i['code'], ('', ''))[1]}"
        + (f" ({i['field']})" if i.get("field") else "")
        for i in row.issues
    )


async def build(session: AsyncSession, batch_id: int) -> bytes:
    book = Workbook()
    book.remove(book.active)
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F4E79")
    used: set[str] = set()
    sheets = await session.scalars(
        select(ImportSheet)
        .where(ImportSheet.batch_id == batch_id)
        .order_by(ImportSheet.number)
    )
    for sheet in sheets:
        out = book.create_sheet(_title(sheet, used))
        kind = KINDS.get(sheet.kind or "")
        columns = sorted(sheet.mapping.items(), key=lambda item: int(item[0]))
        labels = [
            (kind.field(key).label if kind and kind.field(key) else key)
            for _, key in columns
        ]
        out.append(["Строка", "Статус", "Итог", "Код итога", "Замечания", *labels])
        for cell in out[1]:
            cell.font, cell.fill = head, fill
        rows = await session.scalars(
            select(ImportRow)
            .where(ImportRow.sheet_id == sheet.id)
            .order_by(ImportRow.number)
        )
        for row in rows:
            raw = row.data.get("raw", [])
            out.append(
                [
                    row.number,
                    STATUSES.get(row.status, row.status),
                    RESULTS.get(row.result or "", ""),
                    row.result_code or "",
                    _issues(row),
                    *[raw[int(i)] if int(i) < len(raw) else "" for i, _ in columns],
                ]
            )
        for column, width in zip("ABCDE", (8, 16, 14, 11, 50), strict=True):
            out.column_dimensions[column].width = width
        out.freeze_panes = "A2"
    if not book.worksheets:
        book.create_sheet("Итог")
    data = io.BytesIO()
    book.save(data)
    return data.getvalue()
