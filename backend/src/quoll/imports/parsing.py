"""Чтение файла в листы из ячеек-строк (import-design §15.1).

Каждая ячейка сразу приводится к строке: разбор и нормализация дальше
работают с текстом одинаково для всех форматов
"""

import csv
import io
import json
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import openpyxl
import xlrd

from quoll.config import settings
from quoll.imports.errors import import_error, unreadable
from quoll.imports.models import FileFormat

SUFFIXES = {
    ".xlsx": FileFormat.XLSX,
    ".xls": FileFormat.XLS,
    ".csv": FileFormat.CSV,
    ".json": FileFormat.JSON,
}
# заголовок ищется в первых строках листа: выше бывает титул
HEADER_SEARCH_ROWS = 10


@dataclass
class ParsedSheet:
    name: str
    header_row: int = 0
    headers: list[str] = field(default_factory=list)
    # (номер строки в файле с 1, ячейки)
    rows: list[tuple[int, list[str]]] = field(default_factory=list)


def file_format(filename: str | None) -> FileFormat:
    fmt = SUFFIXES.get(Path(filename or "").suffix.lower())
    if fmt is None:
        raise import_error(415, "IMP-002", "Import accepts xlsx, xls, csv or json")
    return fmt


def cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "да" if value else "нет"
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


class _Counter:
    """лимит строк считается по ходу чтения: раздутая книга не дочитывается"""

    def __init__(self):
        self.count = 0

    def add(self) -> None:
        self.count += 1
        if self.count > settings.import_max_rows:
            raise import_error(
                422,
                "IMP-009",
                "Too many rows",
                limit=settings.import_max_rows,
            )


def read(content: bytes, fmt: FileFormat, filename: str) -> list[ParsedSheet]:
    counter = _Counter()
    try:
        raw = {
            FileFormat.XLSX: _xlsx,
            FileFormat.XLS: _xls,
            FileFormat.CSV: _csv,
            FileFormat.JSON: _json,
        }[fmt](content, filename)
        return [_framed(name, rows, counter) for name, rows in raw]
    except Exception as err:
        if getattr(err, "code", None) == "IMP-009":
            raise
        raise unreadable() from None


def _xlsx(content: bytes, filename: str):
    book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        for sheet in book.worksheets:
            yield (
                sheet.title,
                ([cell(v) for v in r] for r in sheet.iter_rows(values_only=True)),
            )
    finally:
        book.close()


def _xls(content: bytes, filename: str):
    book = xlrd.open_workbook(file_contents=content)
    for sheet in book.sheets():

        def rows(sheet=sheet):
            for r in range(sheet.nrows):
                out = []
                for c in range(sheet.ncols):
                    item = sheet.cell(r, c)
                    if item.ctype == xlrd.XL_CELL_DATE:
                        out.append(
                            xlrd.xldate_as_datetime(item.value, book.datemode)
                            .date()
                            .isoformat()
                        )
                    else:
                        out.append(cell(item.value))
                yield out

        yield sheet.name, rows()


def _csv(content: bytes, filename: str):
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("unknown encoding")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    yield Path(filename).stem, ([cell(v) for v in r] for r in reader)


def _json(content: bytes, filename: str):
    data = json.loads(content.decode("utf-8-sig"))
    sheets = (
        {Path(filename).stem: data} if isinstance(data, list) else dict(data.items())
    )
    for name, items in sheets.items():
        if not isinstance(items, list):
            raise ValueError("sheet is not an array")
        keys: list[str] = []
        for item in items:
            for key in item:
                if key not in keys:
                    keys.append(key)

        def rows(items=items, keys=keys):
            yield keys
            for item in items:
                yield [_json_cell(item.get(k)) for k in keys]

        yield str(name), rows()


def _json_cell(value) -> str:
    if isinstance(value, dict | list):
        return json.dumps(value, ensure_ascii=False)
    return cell(value)


def _framed(name: str, rows, counter: _Counter) -> ParsedSheet:
    """строка заголовков - первая из первых строк, где непустых ячеек не
    меньше min(2, ширина листа по этим строкам); выше - титул, пустые
    строки данных пропускаются"""
    sheet = ParsedSheet(name=name)
    rows = iter(rows)
    head = []
    for number, cells in enumerate(rows, start=1):
        head.append((number, _trim(cells)))
        if number >= HEADER_SEARCH_ROWS:
            break
    width = max((len(c) for _, c in head), default=0)
    found = next(
        (
            i
            for i, (_, cells) in enumerate(head)
            if cells and sum(1 for v in cells if v) >= min(2, width)
        ),
        None,
    )
    if found is None:
        return sheet
    sheet.header_row, sheet.headers = head[found]
    rest = [(n, c) for n, c in head[found + 1 :]]
    tail = ((n, c) for n, c in enumerate(rows, start=len(head) + 1))
    for number, cells in _chain(rest, tail):
        if any(cells):
            counter.add()
            sheet.rows.append((number, _fit(cells, len(sheet.headers))))
    return sheet


def _chain(first, second):
    yield from first
    yield from second


def _trim(cells: list[str]) -> list[str]:
    while cells and not cells[-1]:
        cells = cells[:-1]
    return cells


def _fit(cells: list[str], width: int) -> list[str]:
    return (cells + [""] * width)[:width]
