"""Описание формата из spec (В20): ручка видов, шаблон xlsx с «Инструкцией»,
docs/import-format.md. Документ пересобирается командой

    python -m quoll.imports.template docs
"""

import io
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

from quoll.imports.spec import KINDS, FieldSpec, FieldType, KindSpec

TYPES = {
    FieldType.TEXT: "текст",
    FieldType.INT: "целое число",
    FieldType.DATE: "дата: ДД.ММ.ГГГГ, ДД.ММ.ГГ, ДД/ММ/ГГГГ или ГГГГ-ММ-ДД",
    FieldType.BOOL: "да / нет",
    FieldType.LIST: "список через «;»",
    FieldType.REF: "название записи справочника",
    FieldType.LIST_REF: "названия записей справочника через «;»",
    FieldType.ENUM: "одно из значений",
    FieldType.INN: "ИНН: 10 цифр",
    FieldType.KPP: "КПП: 9 цифр",
    FieldType.PHONE: "телефон, приводится к +7…",
    FieldType.EMAIL: "email",
    FieldType.URL: "адрес сайта",
    FieldType.OKSO: "код ОКСО вида 09.03.01",
    FieldType.REGION: "субъект РФ, как в справочнике регионов",
}

RULES = (
    "Файл: xlsx, xls, csv или json, до 10 МБ и до 10 000 строк.",
    "Лист - один вид данных. Вид определяется по имени листа; у файла из "
    "одного листа - по заголовкам. Вид и колонки можно поправить в превью.",
    "Строка заголовков - первая из первых 10 строк, где заполнено хотя бы две "
    "ячейки. «*» у обязательных колонок можно оставить.",
    "Колонки паспорта, СНИЛС, даты рождения и адресов не загружаются.",
    "Пустая ячейка ничего не стирает. Списки-связи только пополняются.",
    "Строки сопоставляются с тем, что уже есть, по ключу: название; продукт - "
    "название и вендор; специальность - код; вуз - ИНН и КПП; контакт - "
    "владелец и ФИО. Повтор ключа в файле сливается, противоречие - ошибка.",
    "Строка может ссылаться на запись, которую создаёт другая строка файла.",
    "Импорт ничего не удаляет. Повторная загрузка того же файла ничего не меняет.",
    "Строка примера в шаблоне - образец: удалите или замените её.",
)


def kinds() -> list[dict]:
    return [
        {
            "key": k.key,
            "label": k.label,
            "order": k.order,
            "fields": [_field(f) for f in k.fields],
        }
        for k in sorted(KINDS.values(), key=lambda k: k.order)
    ]


def _field(f: FieldSpec) -> dict:
    return {
        "key": f.key,
        "label": f.label,
        "type": f.type.value,
        "required": f.required,
        "contract": f.contract,
        "group": f.group,
        "extra": f.extra,
        "ignored": f.ignored,
        "description": f.description,
        "example": f.example,
        "synonyms": list(f.synonyms),
        "enum_values": dict(f.enum_values),
    }


def _format(f: FieldSpec) -> str:
    out = TYPES[f.type]
    if f.enum_values:
        out += ": " + ", ".join(dict.fromkeys(f.enum_values))
    if f.description:
        out += f"; {f.description}"
    if f.contract:
        out += "; у существующих записей не перезаписывается"
    if f.ignored:
        out += "; распознаётся, но не загружается"
    return out


def _header(f: FieldSpec) -> str:
    return f"{f.label}*" if f.required else f.label


def xlsx(keys: list[str] | None = None) -> bytes:
    """лист на вид: заголовки и строка примера; «Инструкция» - первым"""
    chosen = [
        k
        for k in sorted(KINDS.values(), key=lambda k: k.order)
        if not keys or k.key in keys
    ]
    book = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F4E79")
    muted = Font(italic=True, color="808080")
    guide = book.active
    guide.title = "Инструкция"
    guide.append(["Импорт данных в Quoll"])
    guide["A1"].font = Font(bold=True, size=14)
    for rule in RULES:
        guide.append([rule])
    guide.append([])
    guide.append(
        ["Лист", "Колонка", "Обязательная", "Формат", "Пример", "Другие названия"]
    )
    for cell in guide[guide.max_row]:
        cell.font, cell.fill = head, fill
    for kind in chosen:
        for f in kind.fields:
            guide.append(
                [kind.label, f.label, "да" if f.required else "", _format(f), f.example,
                 ", ".join(f.synonyms)]
            )  # fmt: skip
    for column, width in zip("ABCDEF", (24, 30, 14, 60, 30, 40), strict=True):
        guide.column_dimensions[column].width = width
    for row in guide.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for kind in chosen:
        sheet = book.create_sheet(kind.label[:31])
        fields = [f for f in kind.fields if not f.ignored]
        sheet.append([_header(f) for f in fields])
        sheet.append([f.example for f in fields])
        for i, f in enumerate(fields, start=1):
            cell = sheet.cell(1, i)
            cell.font, cell.fill = head, fill
            cell.comment = Comment(_format(f), "Quoll")
            sheet.cell(2, i).font = muted
            sheet.column_dimensions[cell.column_letter].width = max(
                16, len(f.label) + 4
            )
        sheet.freeze_panes = "A2"
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def markdown() -> str:
    lines = [
        "# Формат импорта",
        "",
        "Документ собран из `backend/src/quoll/imports/spec.py` командой "
        "`python -m quoll.imports.template docs`; руками не правится.",
        "",
        "## Общие правила",
        "",
        *[f"- {r}" for r in RULES],
        "",
        "## Листы и колонки",
        "",
        "Виды применяются по порядку: " + " → ".join(k.label for k in _ordered()) + ".",
    ]
    for kind in _ordered():
        lines += [
            "",
            f"### {kind.order}. {kind.label}",
            "",
            "Имя листа: " + ", ".join(f"«{s}»" for s in kind.sheet_synonyms) + ".",
            "",
            "| Колонка | Обяз. | Формат | Пример | Другие названия |",
            "|---|---|---|---|---|",
        ]
        for f in kind.fields:
            lines.append(
                f"| {f.label} | {'да' if f.required else ''} | {_md(_format(f))} "
                f"| {_md(f.example)} | {_md(', '.join(f.synonyms))} |"
            )
    return "\n".join(lines) + "\n"


def _ordered() -> list[KindSpec]:
    return sorted(KINDS.values(), key=lambda k: k.order)


def _md(text: str) -> str:
    return text.replace("|", "\\|")


if __name__ == "__main__" and sys.argv[1:] == ["docs"]:
    target = Path(__file__).resolve().parents[4] / "docs" / "import-format.md"
    target.write_text(markdown(), encoding="utf-8")
    print(f"written {target}")
